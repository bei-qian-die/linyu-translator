# -*- coding: utf-8 -*-
"""生成《淋语翻译器 · 微信小程序端 开发计划书》Word 文档（基于现有 v1.13.0 代码库）。"""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

EA_FONT = "微软雅黑"
ASCII_FONT = "Calibri"
MONO_FONT = "Consolas"

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = ASCII_FONT
normal.font.size = Pt(10.5)
rpr = normal.element.get_or_add_rPr()
rfonts = rpr.get_or_add_rFonts()
rfonts.set(qn("w:eastAsia"), EA_FONT)


def set_ea(run):
    run.font.name = ASCII_FONT
    r = run._element.get_or_add_rPr()
    rf = r.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), EA_FONT)


def H(text, level):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_ea(run)
    return h


def P(text="", bold=False, mono=False, size=None, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        run.bold = bold
        if mono:
            run.font.name = MONO_FONT
            run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA_FONT)
        else:
            set_ea(run)
        if size:
            run.font.size = Pt(size)
    return p


def bullets(items):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(it)
        set_ea(run)


def code(text):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line)
        run.font.name = MONO_FONT
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), EA_FONT)
        run.font.size = Pt(9.5)


def table(headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        set_ea(run)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(val)
            set_ea(run)
    return t


# ============ 标题区 ============
title = doc.add_heading("淋语翻译器 · 微信小程序端 开发计划书", level=0)
for run in title.runs:
    set_ea(run)
P("版本：v1.0（草案）    编制日期：2026-10-21    基线：现有项目 v1.13.0", align=WD_ALIGN_PARAGRAPH.CENTER)
P("")

# ============ 一、项目概述 ============
H("一、项目概述", 1)
H("1.1 现状盘点（基于现有代码库）", 2)
P("「淋语翻译器」已发布 PC 桌面端与安卓端，本项目代码库（linyu-translator v1.13.0）现状如下：")
table(
    ["端", "技术栈", "数据存储", "状态"],
    [
        ["PC 桌面", "Electron 33 + Vue 3 + sql.js", "%APPDATA%\\linyu-translator\\linyu.db", "已发布 v1.13.0"],
        ["Android", "Capacitor 6 套壳，复用 renderer/core", "App 内 files/linyu.db（sql.js）", "已发布 v1.13.0"],
        ["共享核心", "src/core：db / auth / lore / tokenizer / translator", "seed.json + lore-extra3..7.js", "两套端共用"],
        ["微信小程序", "待定（本计划书）", "只读 JSON + wx.storage", "待开发"],
    ],
)
P("内容规模：词库 535 / 句式 146 / 词卡 422 / 人物传 73 / 名言 462。")
P("翻译引擎：src/core/translator.js + tokenizer.js，纯本地「分词 + 正则/词典替换」，无网络依赖；已在 Capacitor WebView（非 Node 的浏览器沙箱）中跑通，说明核心逻辑是平台无关的纯 JS。")

H("1.2 小程序端定位", 2)
bullets([
    "一句话定位：微信里随手可用的「淋语翻译」小工具。",
    "目标用户：喜欢网络梗、追星文化、乐于社交分享的年轻用户。",
    "使用场景：聊天前把话「翻译」成淋语发出去，或生成好玩文案分享到群 / 朋友圈。",
])

H("1.3 项目目标", 2)
bullets([
    "在 2–3 周内完成首版并提审上线。",
    "翻译结果与桌面 / 安卓端保持一致（共用同一份 seed + lore 词库）。",
    "首版实现：输入翻译、一键复制、历史记录、一键分享。",
])

# ============ 二、目标与范围 ============
H("二、目标与范围", 1)
H("2.1 首版功能（MVP+）", 2)
table(
    ["模块", "功能", "说明"],
    [
        ["输入", "文本输入 / 粘贴", "支持多行，设置字数上限（建议 500 字）"],
        ["翻译", "实时 / 点击翻译", "本地规则引擎即时输出"],
        ["结果", "结果展示与复制", "一键复制到剪贴板，成功提示"],
        ["历史", "历史记录", "本地存储最近 N 条，支持删除与清空"],
        ["分享", "分享好友 / 朋友圈", "小程序卡片携带翻译结果文案"],
        ["关于", "使用说明 / 反馈", "版本信息、使用引导、反馈入口"],
    ],
)

H("2.2 明确不在首版范围（把桌面 / 移动的重活砍掉）", 2)
P("小程序首版定位为「面向普通用户的翻译工具」，以下桌面 / 移动专属能力不下放：")
bullets([
    "开发者模式登录（scrypt）与五大板块 CRUD 增删改。",
    "导入 / 导出 JSON、备份恢复、媒体上传（media: 图片）。",
    "sql.js 数据库（见 3.2，小程序无法使用）。",
    "语音朗读、收藏夹、云同步、每日一句、社区、广告与会员。",
])

H("2.3 首版带哪些内容数据", 2)
bullets([
    "必带：词库 words（535）+ 句式 rules（146），驱动翻译。",
    "可选（建议带，数据很小）：词卡 cards / 名言 quotes / 人物传 heroes 的纯文本只读浏览，不含图片。",
    "不带：人物图（8.41MB）、media: 媒体文件、开发者 CRUD。",
])

# ============ 三、关键架构决策 ============
H("三、关键架构决策（重点）", 1)
P("小程序与前两端的本质差异，集中在「渲染层」和「数据层」两处，翻译引擎则可直接复用。逐层对照：")
table(
    ["层", "桌面 / 移动", "小程序", "结论"],
    [
        ["翻译引擎", "translator.js / tokenizer.js（纯 JS）", "直接移植为小程序模块", "可复用"],
        ["内容数据", "seed.json + lore-extra*.js", "只读加载进内存", "可复用（文本约 0.4MB）"],
        ["渲染层", "Vue 3（DOM，无构建工具）", "WXML / WXSS / JS", "需重写（或 Taro/uni-app）"],
        ["数据层", "sql.js（WASM SQLite + 文件落盘）", "无 Node/fs/WASM 同步落盘", "改用 JSON 内存 + wx.storage"],
        ["桥接能力", "Electron IPC / Capacitor 插件", "微信小程序 API", "按微信 API 重实现"],
        ["auth / CRUD / 导入导出 / 备份 / 媒体", "桌面 / 移动专属", "首版不下放", "砍掉"],
    ],
)

H("3.1 翻译引擎可复用（已确认）", 2)
P("翻译核心是确定性文本变换，本地执行、无网络依赖。translator.js 与 tokenizer.js 是平台无关的纯 JS，可直接移植进小程序，配合同一份 seed + lore 数据，保证三端翻译结果一致。移植后用现有 test/core.test.js 的用例做结果比对。")

H("3.2 必须替换 sql.js（最大技术风险）", 2)
P("小程序 JS 环境没有 Node 的 require/文件系统，也难加载 sql.js 的 WASM 并做同步落盘。因此首版彻底去掉 SQLite：")
bullets([
    "只读内容（词库/句式/词卡/名言/人物）在启动时一次性载入内存的普通对象，随小程序版本发布，无需持久化写入。",
    "历史记录、偏好写入 wx.setStorageSync（单用户上限 10MB，历史场景足够）。",
    "开发者 CRUD 不进入小程序，内容维护仍放在桌面端完成后再随版本分发。",
])

H("3.3 渲染层不可复用（Vue → WXML）", 2)
P("现有渲染层是 Vue 3 + 内联模板 + DOM，微信小程序用 WXML/WXSS 渲染、无 DOM，无法直接运行。两条路：")
bullets([
    "方案 A（首版推荐）：原生小程序（WXML/WXSS/JS）重写 UI。首版仅翻译 / 历史两个主页面，工作量小，且不引入新构建工具链，与项目当前「无构建」的风格一致。",
    "方案 B（战略备选）：Taro / uni-app（Vue 3）重写，可复用 Vue 技能、为未来多端与 TypeScript 化铺路（与 IMPROVE-PLAN 的 P9/P10 方向一致），但会引入 Node 构建工具链。",
])

H("3.4 包体积与存储约束（硬约束）", 2)
bullets([
    "小程序主包 ≤ 2MB；使用分包时总分包 ≤ 20MB。",
    "翻译核心数据（seed.json 60.7KB + lore ~323KB）合计约 0.4MB 文本，完全可入主包。",
    "人物图 71 张共 8.41MB，绝不能进包：首版不上人物图，或用云开发云存储 / CDN 按需拉取。",
    "wx.storage 单用户上限 10MB，只放历史记录，不要把整份词库塞进 storage。",
])

# ============ 四、技术选型 ============
H("四、技术选型", 1)
table(
    ["项目", "推荐", "说明"],
    [
        ["框架", "原生小程序（WXML/WXSS/JS）", "首版推荐；核心纯 JS 直接移植，无新构建链"],
        ["备选框架", "Taro / uni-app（Vue 3）", "想复用 Vue 技能 / 未来多端则选，引入构建工具链"],
        ["语言", "JavaScript（与现有代码一致）", "后续可渐进 TypeScript（对齐 IMPROVE-PLAN P9）"],
        ["翻译引擎", "移植 translator.js + tokenizer.js", "平台无关纯逻辑"],
        ["内容数据", "内存加载 seed.json + lore", "只读，随版本发布"],
        ["历史记录", "wx.setStorageSync / getStorageSync", "上限 10MB，够用"],
        ["分享", "onShareAppMessage / onShareTimeline", "携带翻译结果文案"],
        ["图片（可选）", "云开发云存储 / CDN", "人物图按需外置"],
        ["发布", "微信开发者工具 + miniprogram-ci", "自动上传 / 预览 / 提审"],
    ],
)

# ============ 五、页面与模块 ============
H("五、页面与模块设计", 1)
H("5.1 页面结构", 2)
code("""pages/index/index     翻译主页：输入 + 输出 + 复制 + 分享
pages/history/history  历史记录
pages/lexicon/lexicon  词库 / 句式只读浏览（可选）
pages/about/about      关于 / 反馈""")

H("5.2 交互流程", 2)
P("输入文本 → 触发翻译 → 展示结果（可复制）→ 可分享 / 自动存历史 → 历史页可回看并再次复制。")

H("5.3 关键交互细节", 2)
bullets([
    "实时翻译对输入做 300ms 防抖，避免频繁触发。",
    "空输入 / 超长输入给出校验提示。",
    "复制成功后 toast 轻提示。",
    "分享文案自动携带翻译结果，提升传播效果。",
])

# ============ 六、数据结构设计 ============
H("六、数据结构设计（沿用现有 schema）", 1)
H("6.1 内容数据（只读，沿用现有字段）", 2)
code("""// words（词库，驱动翻译）
{ "zh": "好", "linyu": "惹", "alias": "...", "sense": "...",
  "source": "...", "pop": "高", "rel": [...] }

// rules（句式）
{ "pattern": "正则", "template": "替换模板", "note": "...",
  "scene": "...", "examples": "...", "src": "..." }""")

H("6.2 历史记录（wx.storage）", 2)
code("""{
  "id": "1750000000000",
  "input": "今天天气不错",
  "output": "今天天气不错惹噜",
  "createdAt": 1750000000000
}""")

# ============ 七、开发排期 ============
H("七、开发排期（里程碑）", 1)
table(
    ["阶段", "主要工作", "预计工期", "交付物"],
    [
        ["阶段 0：准备", "注册小程序账号 / 主体认证、ICP 备案、类目、开发工具与环境", "1–2 天", "可用的 AppID 与开发环境"],
        ["阶段 1：核心移植", "移植 translator/tokenizer + seed/lore 为小程序模块，跑 core 用例比对", "2–3 天", "小程序内可调用的翻译模块"],
        ["阶段 2：页面开发", "翻译主页、历史页、分享、词库浏览（可选）", "3–5 天", "可运行的首版小程序"],
        ["阶段 3：体验打磨", "UI 适配、异常处理、加载与性能优化、真机调试", "2–3 天", "体验完善的版本"],
        ["阶段 4：测试与提审", "真机测试、隐私合规配置、内容合规、提审跟进", "2–4 天", "通过审核并上线"],
    ],
)
P("合计：约 2–3 周（可按人力并行压缩）。")

# ============ 八、合规与审核要点 ============
H("八、合规与审核要点（重点）", 1)
bullets([
    "类目选择：工具 > 效率（或「翻译」相关类目），通常无需特殊资质。",
    "隐私合规：配置《用户隐私保护指引》；历史记录仅存本地、不对外传输，需如实说明。",
    "内容安全：涉及网络梗语言，需确保词库不含低俗 / 违法违规内容；接入分享等开放能力时遵守《微信公众平台运营规范》。",
    "ICP 备案：小程序主体需完成 ICP 备案（微信已强制要求）。",
    "名称与类目一致：小程序名称与功能一致，避免误导用户。",
])

# ============ 九、风险与应对 ============
H("九、风险与应对", 1)
table(
    ["风险", "影响", "应对措施"],
    [
        ["sql.js 无法在小程序使用", "数据层需重做", "只读内容内存 JSON + 历史 wx.storage，首版不做 CRUD"],
        ["渲染层 Vue 无法复用", "需重写 UI", "首版仅翻译 / 历史两个页面，量小；或改用 Taro/uni-app"],
        ["主包超 2MB", "无法提交", "核心数据仅 0.4MB；人物图外置或首版不上"],
        ["人物图 8.41MB", "体积爆", "云存储 / CDN 按需加载，或首版不上人物图"],
        ["词库更新依赖发版", "迭代慢", "后续云开发远程词库热更新"],
        ["内容合规（网络梗语言）", "被拒 / 下架", "词库合规筛查、隐私指引、内容兜底过滤"],
        ["历史记录隐私", "用户投诉", "仅存本地、提供一键清空"],
    ],
)

# ============ 十、验收标准 ============
H("十、验收标准", 1)
bullets([
    "翻译结果与桌面 / 安卓一致（同一份 seed + lore，跑 translator 用例比对）。",
    "输入 → 翻译 → 复制 → 历史 → 分享全链路可用。",
    "主包体积 < 2MB；本地翻译耗时 < 100ms。",
    "主流机型与微信版本真机测试通过。",
    "通过微信审核并成功上线。",
])

# ============ 十一、后续迭代路线图 ============
H("十一、后续迭代路线图", 1)
bullets([
    "v1.1：词库 / 句式只读浏览、语音朗读、分享卡片美化。",
    "v1.2：远程词库热更新、历史云同步、数据统计。",
    "v1.3：每日一句、社区互动。",
    "商业化：激励广告、会员（去广告 / 高级功能）。",
])

out = "淋语翻译器小程序开发计划书.docx"
doc.save(out)
print("saved:", out)
